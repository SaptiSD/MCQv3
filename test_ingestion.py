"""Reading question banks people actually have. Run with: python test_ingestion.py

The upload used to understand one layout -- numbered questions, lettered options,
an answer key at the end -- and anything else came back as "Questions ready: 0"
with no word about why. The file that showed it was a Moodle quiz pasted into
Word: unnumbered questions, "Select one:", the right option ticked with a Font
Awesome check, and the balance sheet the questions were about in three tables
that the reader never looked at. These pin the reader that replaced it.
"""

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import grading
import ingestion
import teacher_portal as TP
import ui

FAILURES = []


def show(text):
    """The Windows console is cp1252; keep the report readable regardless."""
    print(str(text).encode("ascii", "backslashreplace").decode("ascii"))


def check(label, actual, expected):
    ok = actual == expected
    if not ok:
        FAILURES.append(label)
    show(("  pass " if ok else "  FAIL ") + label + ("" if ok else f"  expected {expected!r} got {actual!r}"))


class FakeUpload:
    def __init__(self, name, data):
        self.name, self._data = name, data

    def getvalue(self):
        return self._data


def docx_upload(build):
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return FakeUpload("bank.docx", buffer.getvalue())


def read_docx(build):
    return ingestion.read_blocks(ingestion.upload_blocks(docx_upload(build)))


def read_text(raw):
    return ingestion.read_blocks(ingestion.text_blocks(raw))


def option(document, letter, text, *, tick=False, bold=False, highlight=False):
    """An option the way a Moodle review page pastes: the letter in its own plain run."""
    paragraph = document.add_paragraph()
    paragraph.add_run(letter)
    run = paragraph.add_run(f". {text}")
    run.bold = bold or None
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if tick:
        mark = paragraph.add_run("\uf00c ")
        mark.bold = bold or None


STATEMENT = [
    ["WILMINGTON CORPORATION", "", ""],
    ["", "Dec 31. Year 7", "Dec. 31 Year 6"],
    ["Cash and cash equivalents", "$634,527", "$335,597"],
    ["", "", ""],
    ["Total Assets", "$3,647,980", "$2,434,430"],
]


def quiz_four(document):
    document.add_paragraph("Selected financial data for Wilmington Corporation is presented below.")
    table = document.add_table(rows=len(STATEMENT), cols=3)
    for r, row in enumerate(STATEMENT):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("What was Wilmington Corporation\u2019s net operating profit margin (NOPM) in Year 7?")
    document.add_paragraph("Select one: ")
    option(document, "a", "13.8%", tick=True, bold=True, highlight=True)
    option(document, "b", "31.0%")
    option(document, "c", "29.9%")
    option(document, "d", "14.6%")
    document.add_paragraph("Using same data as above, the total debt to equity ratio in Year 7 was:")
    document.add_paragraph("Select one: ")
    # No tick on this one: the highlight alone has to carry it.
    option(document, "a", "0.55", bold=True, highlight=True)
    option(document, "b", "1.07")
    option(document, "c", "0.46")
    option(document, "d", "0.26")


print("== the Moodle quiz pasted into Word ==")
reading = read_docx(quiz_four)
questions = reading["questions"]
check("both unnumbered questions are found", len(questions), 2)
check("question text loses the 'Select one:' furniture",
      questions[0]["question_text"], "What was Wilmington Corporation\u2019s net operating profit margin (NOPM) in Year 7?")
check("lower-case letters are read as options",
      questions[0]["options"], [("A", "13.8%"), ("B", "31.0%"), ("C", "29.9%"), ("D", "14.6%")])
check("the tick is not left in the option text", "\uf00c" in questions[0]["options"][0][1], False)
check("the ticked option is the answer", questions[0]["correct_label"], "A")
check("and says where that came from", questions[0]["key_source"], "\u2713 mark")
check("a highlight alone carries the answer", (questions[1]["correct_label"], questions[1]["key_source"]), ("A", "highlight"))
check("both are multiple choice", {q["question_type"] for q in questions}, {"Multiple choice"})
check("nothing to report", reading["notes"], [])
case = reading["case_material"]
check("the introducing sentence is case material", case.startswith("Selected financial data"), True)
check("the table is case material, numbers exact", "| Cash and cash equivalents | $634,527 | $335,597 |" in case, True)
check("the blank spacer row survives", "|  |  |  |" in case, True)
check("figures are right-aligned", "| --- | ---: | ---: |" in case, True)
check("the strict reading agrees", len(ingestion.parse_report("\n".join([
    "What is 2+2?", "a. 3", "b. 4 \u2713"]))[0]), 1)


print()
print("== a file with only the data in it ==")


def data_only(document):
    document.add_paragraph("Selected financial data for Wilmington Corporation is presented below.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Net sales", "$4,885,340"
    table.cell(1, 0).text, table.cell(1, 1).text = "Net income", "$637,664"


reading = read_docx(data_only)
check("no questions", reading["questions"], [])
check("and it says so, rather than 'Questions ready: 0'", reading["notes"][0].startswith("No questions found."), True)
check("naming what it did find", "1 table" in reading["notes"][0], True)
check("which is kept as case material", "| Net sales | $4,885,340 |" in reading["case_material"], True)


print()
print("== a tick drawn from a symbol font ==")


def wingdings(document):
    document.add_paragraph("Which is a mammal?")
    document.add_paragraph("A) Shark")
    paragraph = document.add_paragraph("B) Whale ")
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings")
    symbol.set(qn("w:char"), "F0FC")
    paragraph.runs[0]._r.append(symbol)


reading = read_docx(wingdings)
check("a Wingdings tick marks the answer", reading["questions"][0]["correct_label"], "B")


print()
print("== formatting only counts when it tells options apart ==")


def all_bold(document):
    document.add_paragraph("Pick one")
    for letter, text in (("A", "x"), ("B", "y")):
        document.add_paragraph().add_run(f"{letter}) {text}").bold = True


reading = read_docx(all_bold)
check("every option bold says nothing", reading["questions"][0]["correct_label"], "")
check("and the teacher is asked to fill it in", any("doesn't show which answer" in n for n in reading["notes"]), True)


print()
print("== plain-text layouts ==")
aiken = read_text("What is 2+2?\nA. 3\nB. 4\nANSWER: B\n\nCapital of France?\nA. Paris\nB. Lyon\nANSWER: A\n")
check("Moodle's Aiken format", [(q["correct_label"], q["key_source"]) for q in aiken["questions"]],
      [("B", "answer line"), ("A", "answer line")])
check("unnumbered questions are numbered in order", [q["number"] for q in aiken["questions"]], [1, 2])

moodle = read_text("\n".join([
    "Question 1", "Not yet answered", "Marked out of 1.00", "Flag question", "Question text",
    "Which city is the capital of France?", "Select one:", "a. Lyon", "b. Paris", "c. Nice",
    "The correct answer is: Paris",
    "Question 2", "Correct", "Mark 1.00 out of 1.00", "2 + 2 =", "Select one:", "a. 4", "b. 5",
    "The correct answer is: 4",
]))
check("a Moodle review page's furniture is skipped", [q["question_text"] for q in moodle["questions"]],
      ["Which city is the capital of France?", "2 + 2 ="])
check("its 'Question N' headers set the numbers", [q["number"] for q in moodle["questions"]], [1, 2])
check("'The correct answer is' is matched to the option by text", [q["correct_label"] for q in moodle["questions"]], ["B", "A"])
check("and nothing is reported", moodle["notes"], [])

starred = read_text("1. Is it?\nA) no\n*B) yes\n2. And this?\nA) maybe (correct)\nB) never\n")
check("an asterisk marks the answer", (starred["questions"][0]["correct_label"], starred["questions"][0]["key_source"]), ("B", "asterisk"))
check("so does '(correct)', which is then dropped from the text",
      (starred["questions"][1]["correct_label"], starred["questions"][1]["options"][0][1]), ("A", "maybe"))

conflict = read_text("1. Pick\nA) one \u2713\nB) two\nAnswer key: 1. B\n")
check("an answer key outranks a tick", conflict["questions"][0]["correct_label"], "B")
check("but the disagreement is reported", any("answer key says B" in n and "A" in n for n in conflict["notes"]), True)

several = read_text("Which are even?\na. 2 \u2713\nb. 3\nc. 4 \u2713\n")
check("two ticks make a select-all question", several["questions"][0]["question_type"], grading.SELECT_ALL_TYPE)
check("with both answers", several["questions"][0]["correct_label"], ["A", "C"])
check("and a note to check it", any("more than one answer" in n for n in several["notes"]), True)

hinted = read_text("Which are prime?\nSelect one or more:\na. 2\nb. 4\nc. 5\nANSWER: A, C\n")
check("'Select one or more' means select-all", hinted["questions"][0]["question_type"], grading.SELECT_ALL_TYPE)
check("with no note, because the file said so", hinted["notes"], [])

truth = read_text("The sky is green.\na. True\nb. False \u2713\n")
check("True / False is recognised", truth["questions"][0]["question_type"], "True / False")
check("and keyed by meaning, not by the file's letter", truth["questions"][0]["correct_label"], "B")
flipped = read_text("Water is wet.\na. False\nb. True \u2713\n")
check("even when the file lists False first", (flipped["questions"][0]["options"], flipped["questions"][0]["correct_label"]),
      ([("A", "True"), ("B", "False")], "A"))


print()
print("== nothing goes missing silently ==")
lenient = read_text("1. One?\nA) a\nB) b\n2. Two?\nA) a\nB) b\nAnswer key: 1. A, 3. B\n")
check("a question with no answer is kept, not dropped", [q["number"] for q in lenient["questions"]], [1, 2])
check("with its answer blank", lenient["questions"][1]["correct_label"], "")
check("and a note asking for it", any(n.startswith("Question 2:") for n in lenient["notes"]), True)

off_the_end = read_text("1. One?\nA) a\nB) b\nAnswer key: 1. E\n")
check("an answer that is not an option is blanked", off_the_end["questions"][0]["correct_label"], "")
check("and named", any("says E" in n for n in off_the_end["notes"]), True)

titled = read_text("Biology quiz, chapter 3\n1. One?\nA) a\nB) b\nAnswer key: 1. A\n")
check("a title line is named, not swallowed", any("Biology quiz" in n for n in titled["notes"]), True)
check("and is not case material", titled["case_material"], "")

PASSAGE = ("The mitochondrion is the site of aerobic respiration in eukaryotic cells. It has a double "
           "membrane, its own DNA, and divides independently of the cell it lives in.")
passage = read_text(PASSAGE + "\n1. Where does aerobic respiration happen?\nA) Mitochondrion\nB) Nucleus\nAnswer key: 1. A\n")
check("a passage before question 1 is case material", passage["case_material"], PASSAGE)

before = read_text(PASSAGE + "\nWhat does the passage say divides independently?\na. The mitochondrion \u2713\nb. The nucleus\n")
check("without numbers, a passage is still told apart from its question",
      (before["case_material"], before["questions"][0]["question_text"]),
      (PASSAGE, "What does the passage say divides independently?"))

word = read_text("Pick one\nA) x\nB) y\nC) z\nANSWER: Cafe\n")
check("a word is never read as a string of letters", word["questions"][0]["correct_label"], "")


print()
print("== what the AI is shown ==")
markdown = ingestion.document_markdown(ingestion.upload_blocks(docx_upload(quiz_four)))
check("the tick is spelled out", "[CHECKMARK]" in markdown, True)
check("bold and highlight are marked", "==**a. 13.8% [CHECKMARK]**==" in markdown, True)
check("the tables are there", "| Cash and cash equivalents | $634,527 | $335,597 |" in markdown, True)


print()
print("== case material on screen and on paper ==")
check("dollar signs are escaped, or Streamlit reads them as maths",
      ui.markdown_source("| Cash | $634,527 | $335,597 |"), "| Cash | \\$634,527 | \\$335,597 |")
check("an already-escaped one is left alone", ui.markdown_source("\\$5"), "\\$5")
blocks = TP._markdown_blocks(read_docx(quiz_four)["case_material"])
check("the export sees a sentence and a table", [kind for kind, _ in blocks], ["text", "table"])
rows, right = blocks[1][1]
check("every row of the table, spacer included", len(rows), len(STATEMENT))
check("figures right-aligned", right, [False, True, True])
quiz = {"title": "ZZTEST export", "case_material": read_docx(quiz_four)["case_material"]}
check("the PDF still renders", TP._render_quiz_pdf(quiz, [], "questions")[:4], b"%PDF")
exported = Document(io.BytesIO(TP._render_quiz_docx(quiz, [])))
check("the DOCX carries the statement as a real table", len(exported.tables), 1)
check("with the figures intact", exported.tables[0].cell(2, 1).text, "$634,527")


print()
print("== the chapter quizzes of October 2026 ==")
# A Moodle review page pasted into Word: one-line headers with the theme's
# "Report an issue" and a flag drawn from an icon font, feedback that gives the
# answer away, and the footer naming the TA who made the copy.
REVIEW_COPY = "\n".join([
    "Quiz 15 - copy - Some TA (ta.person@gmail.com) | Copyright 2025 Cambridge Business Publishers",
    "Answer: $52.900",
    "Question 3 Marked out of 14.29  Report an issue",
    "The gain or loss on the sale of the asset is computed as:",
    "a. Market value less net book value",
    "b. Cost less accumulated depreciation ✓",
    "Question 4 Marked out of 14.29  Report an issue",
    "Aiello, Inc. had the following inventory in its fiscal year, which this question and the next one are both "
    "about, so read it with care before you answer anything at all.",
    "You have correctly selected 1.",
    "Which method was used?",
    "a. FIFO",
    "b. LIFO ✓",
])
reading = ingestion.read_blocks(ingestion.text_blocks(REVIEW_COPY))
texts = [question["question_text"] for question in reading["questions"]]
check("the one-line Moodle header is not glued onto the question", texts[0],
      "The gain or loss on the sale of the asset is computed as:")
check("and neither is the icon font's flag", any("" in text for text in texts), False)
check("both questions are read, keys intact", [q["correct_label"] for q in reading["questions"]], ["B", "B"])
case = reading["case_material"]
check("a passage interrupted by feedback is still kept whole", case.startswith("Aiello, Inc."), True)
check("an answer with no question to attach to never reaches it", "52.900" in case, False)
check("nor does the marking feedback", "correctly selected" in case, False)
check("nor somebody's email address", "@gmail.com" in case, False)
left_out = " ".join(note for note in reading["notes"] if note.startswith("Not part of any question"))
check("each of them is named in the notes instead",
      all(part in left_out for part in ("52.900", "correctly selected", "ta.person")), True)
check("and the teacher is told it was a review page", ingestion.REVIEW_PAGE_NOTE in reading["notes"], True)

check("a tick from the icon font survives the clean-up", ingestion._strip_icons("b. 4 "), "b. 4 ")

leaks = ingestion.answer_leaks("Hasbro reports net income of $312.79 million.\n\nYou have correctly selected 1.\n\n"
                               "Question 4(a) — Correct Answer: $25,000\n\nPreesha (desai@example.com)")
check("answers already in a Description/Context are found", len(leaks), 2)
check("quoting the lines", "Correct Answer: $25,000" in leaks[0] and "correctly selected" in leaks[0], True)
check("and the email address", "desai@example.com" in leaks[1], True)
check("a clean Description/Context raises nothing", ingestion.answer_leaks("| Net sales | $4,885,340 |"), [])
check("the \\ufffd sign is counted where it was lost",
      ingestion.unreadable_characters("Kohl�s and Wal-Mart",
                                      [{"question_text": "HAS� equity", "options": []}])[0][:12],
      "2 characters")

typed = ingestion.read_blocks(ingestion.text_blocks("\n".join([
    "1. Which statement reports equity?", "a. Balance sheet", "b. Cash flow", "Answer: A",
    "2. What is the face value of each Series FF share?",
    "Question 2 — Correct Answer: $25,000",
])))
check("“Question 2 — Correct Answer:” keys the typed question above it",
      [("25,000" in str(q["correct_label"])) for q in typed["questions"]], [False, True])


print()
print("== tables from merged Word cells ==")
merged = ingestion.markdown_table([
    ["", "", "", ""],
    ["Ticker", "", "", "Market Cap ($ mil)"],
    ["Advanced Auto Parts (AAP)", "", "", "11,110"],
    ["", "", "", ""],
    ["AutoZone (AZO)", "", "", "28,440"],
])
lines = merged.splitlines()
check("a blank first row no longer becomes an empty header", lines[0], "| Ticker | Market Cap ($ mil) |")
check("the empty columns a merged cell leaves are dropped", lines[2], "| Advanced Auto Parts (AAP) | 11,110 |")
check("a spacer row inside the table stays", lines[3], "|  |  |")
check("figures stay right-aligned", lines[1], "| --- | ---: |")

print()
if FAILURES:
    print(f"{len(FAILURES)} FAILURE(S): {FAILURES}")
    raise SystemExit(1)
print("All ingestion tests passed.")
