"""Question-bank extraction and validation, independent of Streamlit UI.

A file is read as a list of *blocks* in document order -- lines of text, each
remembering how it was formatted, and tables -- and the reader works out which
lines are questions, which are options, which option is marked correct, and
what is left over. Leftover tables (and the sentence introducing them) become
the quiz's case material: an accounting question that says "using the data
above" is unanswerable without the balance sheet it refers to.

The rule this module keeps above all others: **nothing goes missing silently.**
Anything that is not part of a question is either kept as case material or
named in a note, and a question whose answer the file does not give is kept
with a blank answer for the teacher to fill in rather than dropped.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

import grading

SELECT_ALL_TYPE = grading.SELECT_ALL_TYPE

# What an answer copy uses to tick the right option. U+F00C is Font Awesome's
# check, which is what a Moodle review page turns into when pasted into Word.
CHECKS = "\u2713\u2714\u2705\u2611\uf00c\uf058\uf14a\uf046"
# Crosses mark a wrong choice on the same pages. They are stripped, never read.
CROSSES = "\u2717\u2718\u274c\u2612\uf00d\uf057"
# Word's symbol fonts draw ticks from private code points instead.
_SYMBOL_CHECKS = {"wingdings": {"F0FC", "F0FE"}, "wingdings 2": {"F050", "F052"}}

_QUESTION = re.compile(r"^\s*(\d+)[.)](?:\s+|(?=\D))\s*(.+)$")
# The optional asterisk is a common way of marking the right option in a plain
# text file ("*B) Paris"); "(a)" is a common way of lettering one in Word.
_OPTION = re.compile(r"^\s*(\*\s*)?\(?([A-F])[.)]\s*(.+)$", re.I)
_MARKER = re.compile(r"^\s*answer\s*key\s*:?-?\s*(.*)$", re.I)
# "3. A, C" is a select-all answer; the lookahead keeps "1. Answer" from reading
# as question 1, option A.
_KEY_PAIR = re.compile(r"(\d+)\s*[.):-]\s*([A-F](?:\s*,\s*[A-F])*)(?![A-Za-z])", re.I)
# A per-question answer: Moodle's Aiken format ("ANSWER: B") and the feedback
# line a Moodle review page prints under every question.
# "Question 4(a) -- Correct Answer: $25,000" is how one teacher keyed a typed
# question; the label in front of it is not part of the answer.
_KEY_LINE = re.compile(r"^\s*(?:question\s+\S+\s*[—–:-]+\s*)?(?:answer|ans|correct answer)\s*[:=]\s*(.+?)\s*$", re.I)
_CORRECT_IS = re.compile(r"^\s*the\s+correct\s+answers?\s+(?:is|are)\s*:?\s*(.+?)\s*$", re.I)
# Moodle prints its header on one line, and the theme decides what follows the
# number: "Question 3 Marked out of 14.29 Report an issue" in the copy that taught
# this. Matched only as "Question 3", the rest of that line was glued onto the
# next question's wording and the header itself landed in the Description.
_QUESTION_HEADER = re.compile(
    r"^\s*question\s+(\d+)\s*[:.]?"
    r"(?:\s*(?:marked out of [\d.]+|mark [\d.]+ out of [\d.]+|not yet answered|not answered|answer saved"
    r"|complete|correct|incorrect|partially correct|flag question|remove flag|report an issue)\s*[.,]?)*\s*$",
    re.I,
)
# Page furniture that a copied quiz carries along with its questions.
_CHROME = re.compile(
    r"^\s*(?:select one(?: or more)?|select all that apply|choose (?:one|all that apply|the (?:best|correct) answer)"
    r"|not yet answered|not answered|answer saved|marked out of [\d.]+|mark [\d.]+ out of [\d.]+"
    r"|flag question|remove flag|question text|correct|incorrect|partially correct|feedback|information"
    r"|your answer is (?:correct|incorrect|partially correct)|report an issue)\s*[:.!]?\s*$",
    re.I,
)
# Moodle's icon fonts (the flag beside each question, say) arrive as characters
# from the Private Use Area that no other font can draw. The tick and cross
# glyphs are among them and *are* meaning -- they mark the answer -- so they stay.
_ICON_GLYPHS = re.compile("[-]")
# Lines that give an answer away. In the Description/Context every student reads
# them above every question; a copy of Moodle's *review* page is full of them.
_ANSWER_GIVEAWAY = re.compile(
    r"\b(?:correct answer|the correct answers? (?:is|are)|you have correctly selected|your answer is correct"
    r"|answer\s*[:=]|ans\s*[:=]|solution\s*[:=]|answer key)",
    re.I,
)
_REVIEW_PAGE = re.compile(r"you have correctly selected|your answer is (?:correct|incorrect)", re.I)
# Where a review page's marking feedback starts. Everything after it, up to the
# next question's header, is the worked solution -- "Bava reports $52,900 as
# income: ..." -- and kept as Description/Context it handed every student the
# answer.
_FEEDBACK_START = re.compile(r"^\s*(?:solution|general feedback|feedback|explanation|the correct answer is)\b\s*:?", re.I)
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_SELECT_ALL_HINT = re.compile(r"select (?:one or more|all that apply)|choose all that apply", re.I)
_CORRECT_TAG = re.compile(r"\s*[\(\[]\s*correct\s*[\)\]]\s*$", re.I)
# Separated letters only: "A, C" or "a and c", never "Cafe" read as C, A, F, E.
_LETTERS = re.compile(r"^\(?[A-F]\)?(?:(?:\s*(?:,|and|&)\s*|\s+)\(?[A-F]\)?)*\.?$", re.I)
_NUMERIC_CELL = re.compile(r"^[\s$\u20ac\u00a3(),.%\d\u2212-]+$")

# Prose before the first question is a passage worth keeping only if there is
# enough of it to be one; a title line is not case material.
PASSAGE_CHARACTERS = 150


# -- Reading the file ---------------------------------------------------------


def _decode(raw: bytes) -> str:
    # A byte-order mark says outright what the encoding is, so trust it. Without
    # one, utf-16 must never be guessed at: `bytes.decode("utf-16")` succeeds on
    # almost any even-length input by pairing bytes up, so a Windows ANSI file --
    # what "Save as plain text" writes the moment the document contains a curly
    # apostrophe -- came back as CJK noise and parsed as nothing at all. Because
    # it turned on the file's length being even, adding a space to the end could
    # fix or break it, which made it look intermittent.
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    # utf-8-sig also eats the utf-8 BOM, which is not whitespace to `re` and so
    # kept question 1 from ever matching. cp1252 is the Windows ANSI fallback.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def extract_upload(upload) -> str:
    """The upload as plain text: a .docx's paragraphs, or a text file decoded."""
    if upload.name.lower().endswith(".docx"):
        return "\n".join(p.text for p in Document(io.BytesIO(upload.getvalue())).paragraphs)
    return _decode(upload.getvalue())


def upload_blocks(upload) -> list[dict]:
    """The upload as blocks, in document order.

    Reading a .docx through `Document.paragraphs`, as this used to, skips every
    table without a trace -- the balance sheet a question was about simply
    vanished, and nothing on screen said so.
    """
    if upload.name.lower().endswith(".docx"):
        return _docx_blocks(Document(io.BytesIO(upload.getvalue())))
    return text_blocks(_decode(upload.getvalue()))


def text_blocks(raw: str) -> list[dict]:
    return [{"kind": "text", "text": line, "marks": set()} for line in raw.splitlines()]


def _docx_blocks(document) -> list[dict]:
    blocks: list[dict] = []
    icons: dict = {}
    for element in _body_elements(document.element.body):
        if element.tag == qn("w:p"):
            blocks.extend(_paragraph_lines(Paragraph(element, document), icons))
        else:
            rows = _table_rows(Table(element, document))
            if any(cell for row in rows for cell in row):
                blocks.append({"kind": "table", "rows": rows})
    return blocks


def _body_elements(parent):
    for child in parent.iterchildren():
        if child.tag in (qn("w:p"), qn("w:tbl")):
            yield child
        elif child.tag == qn("w:sdt"):
            # A content control wraps ordinary paragraphs and tables.
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _body_elements(content)


def _paragraph_lines(paragraph, icons: dict | None = None) -> list[dict]:
    """A paragraph as one block per line, each with the marks its text carries.

    A soft line break is a new line to the reader, because a question pasted
    from the web often arrives as one paragraph with its options on breaks.

    Each line also carries its `boxes` -- the text of any text box anchored in
    it -- and its `icons`. Both are kept apart from `text` on purpose: in a
    filled-in answer copy a text box *is* the answer (a printed quiz's answer
    fields become text boxes), and the text goes on to become Description/
    Context, which every student reads. Only the AI reader is shown them.
    """
    icons = {} if icons is None else icons
    lines: list[list[tuple[str, set]]] = [[]]
    extras: list[tuple[list, list]] = [([], [])]
    for item in paragraph.iter_inner_content():
        for run in getattr(item, "runs", [item]):
            marks = _run_marks(run)
            boxes, pictures = _run_extras(run, icons)
            for index, piece in enumerate((run.text + _symbol_checks(run)).split("\n")):
                if index:
                    lines.append([])
                    extras.append(([], []))
                lines[-1].append((piece, marks))
            extras[-1][0].extend(boxes)
            extras[-1][1].extend(pictures)
    return [{"kind": "text", "text": "".join(piece for piece, _ in parts), "marks": _line_marks(parts),
             "boxes": boxes, "icons": pictures}
            for parts, (boxes, pictures) in zip(lines, extras)]


_TEXT_BOX = qn("w:txbxContent")
_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_EMBED = qn("r:embed")
# Bigger than this is a chart or a photo, not a radio button.
ICON_PIXELS = 64


def _inside(element, tag: str) -> bool:
    parent = element.getparent()
    while parent is not None:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _run_extras(run, icons: dict) -> tuple[list[str], list[str]]:
    """The text boxes and small pictures drawn in `run`.

    Word stores a text box twice -- the modern drawing and a VML fallback for
    old readers -- so anything under `mc:Fallback` is the same box again.
    """
    boxes, pictures = [], []
    for element in run._r.iter(_TEXT_BOX, _BLIP):
        if _inside(element, _FALLBACK):
            continue
        if element.tag == _TEXT_BOX:
            text = " ".join("".join(node.text or "" for node in element.iter(qn("w:t"))).split())
            if text:
                boxes.append(text)
        elif (label := _icon(run.part, element.get(_EMBED), icons)):
            pictures.append(label)
    return boxes, pictures


def _icon(part, relationship: str | None, icons: dict) -> str | None:
    """A small picture's label, the same label for pictures that look alike.

    A quiz copied from Moodle marks the chosen option with a filled radio button
    and the rest with empty ones, each saved as a separate image file -- so the
    files differ even where the pictures don't. An 8x8 average-hash says which
    ones look the same, and the option whose picture looks different is the
    marked one.
    """
    try:
        from PIL import Image
        blob = part.related_parts[relationship].blob
        with Image.open(io.BytesIO(blob)) as image:
            if max(image.size) > ICON_PIXELS:
                return None
            small = list(image.convert("L").resize((8, 8)).getdata())
    except Exception:
        return None
    mean = sum(small) / len(small)
    bits = tuple(value > mean for value in small)
    for label, known in icons.items():
        if sum(a != b for a, b in zip(bits, known)) <= 6:
            return label
    label = f"icon {chr(ord('A') + len(icons))}" if len(icons) < 26 else f"icon {len(icons)}"
    icons[label] = bits
    return label


def _symbol_checks(run) -> str:
    ticks = ""
    for symbol in run._r.findall(qn("w:sym")):
        font = (symbol.get(qn("w:font")) or "").strip().casefold()
        char = (symbol.get(qn("w:char")) or "").upper().rjust(4, "0")
        if char in _SYMBOL_CHECKS.get(font, set()):
            ticks += "\u2713"
    return ticks


def _run_marks(run) -> set:
    marks = set()
    if run.bold:
        marks.add("bold")
    if run.font.highlight_color is not None:
        marks.add("highlight")
    else:
        properties = run._r.rPr
        shading = properties.find(qn("w:shd")) if properties is not None else None
        if shading is not None and (shading.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF"):
            marks.add("highlight")
    colour = run.font.color
    if colour is not None and colour.type is not None:
        value = str(colour.rgb or colour.theme_color or "")
        if value and value.upper() not in ("000000", "TEXT_1 (13)", "AUTO"):
            marks.add(f"colour {value}")
    return marks


def _line_marks(parts) -> set:
    """Marks every *meaningful* piece of a line shares.

    The option letter is often its own run in a different style -- Moodle copies
    put "a" plain and ". 13.8%" bold and highlighted -- so pieces that are only
    a letter, punctuation or a tick do not get a vote.
    """
    voting = [marks for piece, marks in parts if _meaningful(piece)]
    if not voting:
        return set()
    return set.intersection(*voting)


def _meaningful(piece: str) -> bool:
    core = re.sub(r"[\W_]+", "", piece)
    return bool(core) and not re.fullmatch(r"[A-Fa-f]", core)


def _table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells, previous = [], None
        for cell in row.cells:
            # A merged cell comes back once per grid column it spans.
            cells.append("" if cell._tc is previous else " ".join(cell.text.split()))
            previous = cell._tc
        rows.append(cells)
    return rows


def markdown_table(rows: list[list[str]]) -> str:
    """A table as Markdown, numeric columns right-aligned like a statement.

    Columns with nothing in them are left out, and so are blank rows before the
    first real one. Word reports a merged cell once per grid column it spans, so
    a company name merged across three columns drew two empty columns beside
    every row; and a blank first row became an empty header that pushed the
    real one into the body. A blank row *inside* a statement is a spacer the
    teacher put there, and stays.
    """
    width = max(len(row) for row in rows)
    rows = [[cell.replace("|", "\\|") for cell in row] + [""] * (width - len(row)) for row in rows]
    while len(rows) > 1 and not any(cell.strip() for cell in rows[0]):
        rows = rows[1:]
    kept = [column for column in range(width) if any(row[column].strip() for row in rows)] or [0]
    rows = [[row[column] for column in kept] for row in rows]
    width = len(kept)
    head, body = rows[0], rows[1:]
    alignment = []
    for column in range(width):
        values = [row[column] for row in body if row[column]]
        numeric = values and sum(bool(_NUMERIC_CELL.match(v)) for v in values) >= 0.6 * len(values)
        alignment.append(" ---: |" if numeric else " --- |")
    lines = ["| " + " | ".join(head) + " |", "|" + "".join(alignment)]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)


_BLANK_CELL = re.compile(r"^[\s$€£_.]*$")


def mark_filled_cells(question_blocks: list[dict], answer_blocks: list[dict]) -> list[dict]:
    """The answer copy's blocks, with every table cell the student filled in
    written as an answer box.

    A fill-in-the-table question -- "Basic earnings per share | a. | $ ___" --
    has its answer in the answer copy's cell and nowhere else, and nothing about
    "$ 2.64" in a table says it is an answer rather than a figure. Its twin in
    the question file is blank, so the difference does: a cell that is empty
    (or just "$") in the question file and holds something in the answer copy
    is what was filled in. Tables are paired in order and only when their shapes
    match; anything else is left exactly as it was.
    """
    question_tables = [block for block in question_blocks if block["kind"] == "table"]
    answer_tables = [block for block in answer_blocks if block["kind"] == "table"]
    if len(question_tables) != len(answer_tables):
        return answer_blocks
    filled = {}
    for question, answer in zip(question_tables, answer_tables):
        if [len(row) for row in question["rows"]] != [len(row) for row in answer["rows"]]:
            continue
        rows = [[f"[answer box: {cell.strip()}]" if _BLANK_CELL.match(blank) and not _BLANK_CELL.match(cell) else cell
                 for blank, cell in zip(blank_row, row)]
                for blank_row, row in zip(question["rows"], answer["rows"])]
        filled[id(answer)] = {**answer, "rows": rows}
    return [filled.get(id(block), block) for block in answer_blocks]


def document_markdown(blocks: list[dict]) -> str:
    """The blocks as Markdown, formatting marks included, for the AI reader.

    The marks are the whole point: a model cannot see that an option was bold
    or highlighted unless it is told, and that is how answer copies say which
    option is right.
    """
    out = []
    for block in blocks:
        if block["kind"] == "table":
            out.append(markdown_table(block["rows"]))
            continue
        line = block["text"].strip()
        boxes = [f"[answer box: {text}]" for text in block.get("boxes", [])]
        pictures = [f"[{label}]" for label in block.get("icons", [])]
        if not line and not boxes:
            continue
        line = " ".join([*pictures, line, *boxes]).strip()
        for tick in CHECKS:
            line = line.replace(tick, " [CHECKMARK] ")
        line = " ".join(line.split())
        marks = block["marks"]
        if "bold" in marks:
            line = f"**{line}**"
        if "highlight" in marks:
            line = f"=={line}=="
        colour = next((mark for mark in marks if mark.startswith("colour ")), None)
        if colour:
            line = f"[{colour}] {line}"
        out.append(line)
    return "\n".join(out)


# -- Making sense of it -------------------------------------------------------


def parse_bank(raw: str) -> list[dict]:
    return parse_report(raw)[0]


def parse_report(raw: str) -> tuple[list[dict], list[str]]:
    """The strict reading: only questions whose answer the file gives, plus a note
    about every question it had to drop.

    Dropping questions silently is how a 40-question upload becomes a
    35-question quiz with nobody the wiser.
    """
    questions, skipped = [], []
    for question in read_blocks(text_blocks(raw))["questions"]:
        labels = {label for label, _ in question["options"]}
        correct = question["correct_label"]
        given = question.get("given_key")
        if question["options"] and correct and set(correct if isinstance(correct, list) else [correct]) <= labels:
            questions.append(question)
        elif given:
            skipped.append(f"Question {question['number']}: the answer key says {given}, which is not one of its options")
        else:
            skipped.append(f"Question {question['number']}: no entry in the answer key")
    return questions, skipped


def read_blocks(blocks: list[dict]) -> dict:
    """Everything the review screen needs from one file.

    Returns `questions` (every one found, with a blank `correct_label` where the
    file does not say), `case_material` (Markdown), and `notes` -- plain
    sentences about anything the teacher has to check or fill in.
    """
    reader = _Reader(numbered=_uses_numbers(blocks))
    for block in blocks:
        reader.feed(block)
    return reader.finish()


def _uses_numbers(blocks: list[dict]) -> bool:
    """Whether this file numbers its questions ("1. ...", followed by options).

    A numbered file is read exactly as it always was. A file without numbers --
    a Moodle quiz pasted into Word -- is read by layout instead: a line of text
    followed by lettered options is a question.
    """
    lines = [block["text"].strip() for block in blocks
             if block["kind"] == "text" and block["text"].strip() and not _CHROME.match(block["text"])]
    return any(_QUESTION.match(line) and _OPTION.match(following)
               for line, following in zip(lines, lines[1:]))


class _Reader:
    def __init__(self, numbered: bool):
        self.numbered = numbered
        self.questions: list[dict] = []
        self.current: dict | None = None
        self.pending: list[str] = []      # prose no question has claimed yet
        self.trailing: list[str] = []     # prose after a numbered question's options
        self.case: list[str] = []
        self.ignored: list[str] = []
        self.answers: dict[int, str] = {}
        self.in_answer_key = False
        self.select_all_hint = False
        self.next_number: int | None = None
        self.stray_options = 0
        self.tables = 0
        self.review_page = False
        self.in_feedback = False

    # Each block, in order.

    def feed(self, block: dict) -> None:
        if block["kind"] == "table":
            self.tables += 1
            if self.in_answer_key:
                for row in block["rows"]:
                    self._answer_line(". ".join(cell for cell in row if cell))
                return
            self._claim_pending_as_case()
            self._claim_trailing_as_case()
            self.case.append(markdown_table(block["rows"]))
            return
        line = " ".join(_strip_icons(block["text"]).split())
        if not line:
            return
        if (marker := _MARKER.match(line)):
            self._settle()
            self.in_answer_key = True
            self._answer_line(marker.group(1))
            return
        if self.in_answer_key:
            self._answer_line(line)
            return
        if _CHROME.match(line):
            if _SELECT_ALL_HINT.search(line):
                if self.current is not None and not self.current["options"] and self.numbered:
                    self.current["select_all"] = True
                else:
                    self.select_all_hint = True
            return
        if (header := _QUESTION_HEADER.match(line)):
            self._settle()
            self.in_feedback = False
            self.next_number = int(header.group(1))
            return
        key = _KEY_LINE.match(line) or _CORRECT_IS.match(line)
        if key and self.current is not None:
            self.current["key_line"] = key.group(1)
            return
        numbered_start = self.numbered and _QUESTION.match(line)
        if self.in_feedback and not numbered_start:
            self.ignored.append(line)
            return
        self.in_feedback = False
        if _REVIEW_PAGE.search(line) or _FEEDBACK_START.match(line):
            self.review_page = True
            self.in_feedback = True
            self._claim_trailing_as_case()
            self.ignored.append(line)
            return
        numbered = _QUESTION.match(line) if self.numbered else None
        option = _OPTION.match(line)
        if numbered:
            self._settle()
            self._start(numbered.group(2), int(numbered.group(1)))
        elif option:
            self._option(option, block["marks"])
        elif self.numbered:
            if self.current is None:
                self._unclaimed(line)
            elif not self.current["options"]:
                self.current["question_text"] += " " + line
            else:
                self.trailing.append(line)
        else:
            self._unclaimed(line)

    def _unclaimed(self, line: str) -> None:
        """Prose no question has claimed: on its way to the Description/Context.

        Unless it is an answer with no question to attach it to, or somebody's
        email address. Every student reads the Description/Context, so those are
        left out instead, and the notes name them.
        """
        if _KEY_LINE.match(line) or _CORRECT_IS.match(line) or _EMAIL.search(line) or _REVIEW_PAGE.search(line):
            # Not `_settle()`: a passage still being gathered for the next
            # question must not be cut short by a stray line in the middle of it.
            self.review_page = self.review_page or bool(_REVIEW_PAGE.search(line))
            self.ignored.append(line)
            return
        self.pending.append(line)

    def _option(self, match, marks: set) -> None:
        if not self.numbered and (self.current is None or self.pending or not self.current["options"]):
            if self.pending:
                passage, text = _split_passage(self.pending)
                self.pending = []
                if passage:
                    self.case.extend(passage)
                self._start(text, self.next_number)
            elif self.current is None:
                self.stray_options += 1
                return
        if self.current is None:
            self.stray_options += 1
            return
        if self.trailing:
            # A numbered question's prose that runs on after its options.
            self.current["question_text"] += " " + " ".join(self.trailing)
            self.trailing = []
        text, inline = _option_marks(match.group(3))
        if match.group(1):
            inline.add("asterisk")
        if text:
            self.current["options"].append((match.group(2).upper(), text))
            self.current["marks"].append(inline | marks)

    def _start(self, text: str, number: int | None) -> None:
        if number is None:
            number = (self.questions[-1]["number"] + 1) if self.questions else 1
        self.current = {"number": number, "question_text": text, "options": [], "marks": [],
                        "select_all": self.select_all_hint}
        self.select_all_hint = False
        self.next_number = None
        self.questions.append(self.current)

    def _answer_line(self, line: str) -> None:
        for number, letters in _KEY_PAIR.findall(line or ""):
            self.answers[int(number)] = letters

    # Leftovers.

    def _settle(self) -> None:
        """Close the question in progress before something new starts."""
        if self.trailing and self.current is not None:
            self.current["question_text"] += " " + " ".join(self.trailing)
        self.trailing = []
        self._claim_pending_as_leftover()

    def _claim_pending_as_case(self) -> None:
        self.case.extend(self.pending)
        self.pending = []

    def _claim_trailing_as_case(self) -> None:
        self.case.extend(self.trailing)
        self.trailing = []

    def _claim_pending_as_leftover(self) -> None:
        if not self.pending:
            return
        if not self.questions and sum(len(line) for line in self.pending) >= PASSAGE_CHARACTERS:
            self.case.extend(self.pending)
        else:
            self.ignored.extend(self.pending)
        self.pending = []

    def finish(self) -> dict:
        self._settle()
        notes: list[str] = []
        questions = [_resolve(question, self.answers, notes) for question in self.questions]
        if self.ignored:
            shown = [f"\u201c{_clip(line)}\u201d" for line in self.ignored[:5]]
            more = f" and {len(self.ignored) - 5} more line(s)" if len(self.ignored) > 5 else ""
            notes.append("Not part of any question, so left out: " + "; ".join(shown) + more + ".")
        if self.stray_options:
            notes.append(f"{self.stray_options} lettered line(s) had no question before them and were left out.")
        if self.review_page:
            notes.append(REVIEW_PAGE_NOTE)
        case = "\n\n".join(self.case).strip()
        notes.extend(answer_leaks(case) + unreadable_characters(case, questions))
        if not questions:
            kept = ""
            if case and self.tables:
                kept = (f" Its {self.tables} table{'s' if self.tables != 1 else ''} and the text around "
                        f"{'them' if self.tables != 1 else 'it'} were kept as the quiz's Description/Context.")
            elif case:
                kept = " Its text was kept as the quiz's Description/Context."
            notes.insert(0, (
                "No questions found. A question is a line of text followed by lettered options "
                "(A. / B. / C. or a) b) c)), with the right one marked by an answer key, a tick, bold or highlight."
                + kept
            ))
        # `unread` is how much of the file nothing could place -- the case for
        # suggesting the AI reader. A file of nothing but tables has none.
        return {"questions": questions, "case_material": case, "notes": notes, "tables": self.tables,
                "unread": len(self.ignored) + self.stray_options}


def _split_passage(pending: list[str]) -> tuple[list[str], str]:
    """Separate a passage from the question that follows it.

    Without numbers the only boundary is layout, so several paragraphs before a
    set of options all read as one question. When the last one is plainly the
    question and the rest is long enough to be a passage, the rest is case
    material instead.
    """
    if (len(pending) > 1 and pending[-1].rstrip().endswith(("?", ":"))
            and sum(len(line) for line in pending[:-1]) >= PASSAGE_CHARACTERS):
        return pending[:-1], pending[-1]
    return [], " ".join(pending)


def _option_marks(text: str) -> tuple[str, set]:
    marks = set()
    if any(tick in text for tick in CHECKS):
        marks.add("check")
    for character in CHECKS + CROSSES:
        text = text.replace(character, " ")
    if _CORRECT_TAG.search(text):
        marks.add("tag")
        text = _CORRECT_TAG.sub("", text)
    if re.search(r"\s\*+\s*$", text):
        marks.add("asterisk")
        text = re.sub(r"\s\*+\s*$", "", text)
    return " ".join(text.split()), marks


_SOURCES = {"check": "\u2713 mark", "asterisk": "asterisk", "tag": "(correct) label"}


def _resolve(question: dict, answers: dict, notes: list[str]) -> dict:
    """Decide a question's type and answer from whatever the file offered."""
    number, options, marks = question["number"], question["options"], question["marks"]
    labels = [label for label, _ in options]
    explicit, source, given = None, "", None
    if question.get("key_line"):
        explicit, given = _key_labels(question["key_line"], options), question["key_line"]
        source = "answer line"
    elif number in answers:
        explicit, given = _key_labels(answers[number], options), answers[number]
        source = "answer key"

    marked, mark_source = [], ""
    strong = [label for label, option_marks in zip(labels, marks) if option_marks & set(_SOURCES)]
    if strong:
        marked = strong
        mark_source = _SOURCES[next(iter(sorted(marks[labels.index(strong[0])] & set(_SOURCES))))]
    else:
        # Formatting only counts when it tells the options apart: a file with
        # every option in bold is saying nothing about which is right.
        for flag, name in (("highlight", "highlight"), ("bold", "bold")):
            flagged = [label for label, option_marks in zip(labels, marks) if flag in option_marks]
            if 0 < len(flagged) < len(labels):
                marked, mark_source = flagged, name
                break
        else:
            colours = [next((m for m in option_marks if m.startswith("colour ")), None) for option_marks in marks]
            coloured = [label for label, colour in zip(labels, colours) if colour]
            if 0 < len(coloured) < len(labels) and len({c for c in colours if c}) == 1:
                marked, mark_source = coloured, "coloured text"

    if not options:
        # A question with no options is a typed one; its answer line is the answer.
        answer = question.get("key_line") or ""
        if not answer:
            notes.append(f"Question {number}: no options and no answer, so it was read as a typed question "
                         "with its answer left blank. Fill it in, or delete the row.")
        return {**_public(question), "question_type": "Short answer", "correct_label": answer,
                "key_source": "answer line" if answer else "", "given_key": given}

    if explicit is not None:
        correct = explicit
        if marked and set(marked) != set(explicit) and set(explicit) <= set(labels):
            notes.append(f"Question {number}: the {source} says {', '.join(explicit)} but the "
                         f"{mark_source} is on {', '.join(marked)}. The {source} was used; check it.")
    else:
        correct, source = marked, mark_source

    if correct and not set(correct) <= set(labels):
        notes.append(f"Question {number}: the {source} says {given}, which is not one of its options. "
                     "Fill in the right one.")
        correct, source = [], ""
    if not correct:
        if explicit is None:
            notes.append(f"Question {number}: the file doesn't show which answer is right. Fill it in.")
        source = ""

    texts = [grading.normalise_text(text) for _, text in options]
    if sorted(texts) == ["false", "true"]:
        by_text = dict(zip(labels, texts))
        key = "A" if correct and by_text[correct[0]] == "true" else "B" if correct else ""
        return {**_public(question), "options": [("A", "True"), ("B", "False")],
                "question_type": "True / False", "correct_label": key, "key_source": source, "given_key": given}
    if len(correct) > 1 or question.get("select_all"):
        if len(correct) > 1 and not question.get("select_all"):
            notes.append(f"Question {number} has more than one answer marked, so it was read as "
                         "\u201cselect all that apply\u201d. Check that is what you meant.")
        return {**_public(question), "question_type": SELECT_ALL_TYPE, "correct_label": correct,
                "key_source": source, "given_key": given}
    return {**_public(question), "question_type": "Multiple choice", "correct_label": correct[0] if correct else "",
            "key_source": source, "given_key": given}


def _public(question: dict) -> dict:
    return {"number": question["number"], "question_text": question["question_text"].strip(),
            "options": list(question["options"])}


def _key_labels(value: str, options: list) -> list[str]:
    """An answer given as letters ("B", "A, C") or as the option's own text."""
    value = value.strip()
    if _LETTERS.match(value):
        return list(dict.fromkeys(re.findall(r"(?<![A-Z])[A-F](?![A-Z])", value.upper())))
    wanted = grading.normalise_text(value)
    return [label for label, text in options if grading.normalise_text(text) == wanted] or [value]


REVIEW_PAGE_NOTE = (
    "This looks like a copy of a quiz's review page -- it has a student's answers and the marking feedback "
    "in it. Lines that gave answers away were left out, but check the questions and the Description/Context "
    "for any answers still showing, or copy the quiz from its preview page instead."
)


def _strip_icons(text: str) -> str:
    return _ICON_GLYPHS.sub(lambda m: m.group() if m.group() in CHECKS + CROSSES else " ", text)


def answer_leaks(case_material: str) -> list[str]:
    """Warnings for anything in the Description/Context that gives answers away.

    The Description/Context is drawn above every question for every student,
    so an "Answer:" line in it is the answer key handed out with the paper.
    Warn rather than delete: a worked example can legitimately say "answer",
    and silently cutting a teacher's text is how the upload used to lose tables.
    """
    warnings = []
    lines = [" ".join(_strip_icons(line).split()) for line in (case_material or "").splitlines()]
    giveaways = [line for line in lines if _ANSWER_GIVEAWAY.search(line)]
    if giveaways:
        shown = "; ".join(f"“{_clip(line)}”" for line in giveaways[:3])
        more = f" and {len(giveaways) - 3} more" if len(giveaways) > 3 else ""
        warnings.append(f"The Description/Context seems to contain answers, and every student sees it: "
                        f"{shown}{more}. Delete them before activating.")
    emails = sorted({match.group() for line in lines for match in _EMAIL.finditer(line)})
    if emails:
        warnings.append("The Description/Context contains an email address (" + ", ".join(emails[:3])
                        + "), which every student would see. Remove it unless that is intended.")
    return warnings


def unreadable_characters(case_material: str, questions=()) -> list[str]:
    """A note for the � sign, which is a character lost before MCQ saw it.

    It appears where a curly apostrophe or a multiplication sign was copied out
    of a page whose own encoding was broken. Nothing can say which character it
    was, so the teacher is told rather than given a guess.
    """
    places = (case_material or "").count("�") + sum(
        str(question.get("question_text", "")).count("�")
        + sum(str(text).count("�") for _, text in question.get("options", []))
        for question in questions)
    if not places:
        return []
    return [f"{places} character{'s' if places != 1 else ''} in this file could not be read and show as � "
            "-- usually an apostrophe or a × sign lost when the text was copied. Retype "
            f"{'them' if places != 1 else 'it'} before publishing."]


def _clip(line: str, limit: int = 70) -> str:
    return line if len(line) <= limit else line[: limit - 1].rstrip() + "\u2026"
